"""Build the printable draft invigilation table (same layout as the hand-made draft) from data.json.

Usage: python3 build_draft_xlsx.py data.json out.xlsx
"""
import json, re, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

data = json.load(open(sys.argv[1], encoding='utf-8'))
out = sys.argv[2]
people = {p['id']: p for p in data['people']}
duties = data['duties']
rooms = data['rooms']
room_bldg = data.get('roomBldg', {})
dist_bldg = data.get('distBldg', {})

DOW = ['จันทร์', 'อังคาร', 'พุธ', 'พฤหัสบดี', 'ศุกร์', 'เสาร์', 'อาทิตย์']
MON = ['มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน', 'กรกฎาคม',
       'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม']
TITLE = re.compile(r'^(นางสาว|นาง|นาย)')
ACAD = re.compile(r'^((?:รศ\.|ผศ\.|ศ\.|อ\.)?(?:ดร\.)?)')


def thai_day(s):
    d = dt.date.fromisoformat(s)
    return f'วัน{DOW[d.weekday()]}ที่ {d.day} {MON[d.month - 1]} {d.year + 543}'


def short(name):
    """Staff -> first name; teacher -> academic title + first name (as in the hand-made draft)."""
    name = re.sub(r'\s*\(.*\)\s*$', '', name or '')
    if TITLE.match(name):
        return TITLE.sub('', name).split()[0]
    title = ACAD.match(name).group(1)
    return title + name[len(title):].split()[0]


def names(pids_or_names):
    seen = []
    for n in pids_or_names:
        n = short(n)
        if n and n not in seen:
            seen.append(n)
    return ', '.join(seen)


def at(date, time, role, room=None):
    return [people[d['pid']]['name'] for d in duties
            if d['date'] == date and d['time'] == time and d['role'] == role and (room is None or d['room'] == room)]


def dist_for(date, time, room):
    b = room_bldg.get(room, 'PS')
    return [n for n in at(date, time, 'dist') if dist_bldg.get(f'{date}|{time}|{n}', b) == b]


openers = {d['date']: people[d['pid']]['name'] for d in duties if d['role'] == 'opener'}

# one row per section x room, like the draft
by_day = {}
for key, entries in rooms.items():
    date, time, room = key.split('|')
    teachers = at(date, time, 'own', room)
    staff = at(date, time, 'proctor', room) + [e['distHelp'] for e in entries if e.get('distHelp')]
    t, s = names(teachers), names(staff)
    proctors = f'{t} - {s}' if t and s else (t or s)
    if not proctors and any((e.get('need') or 0) > 0 for e in entries):
        proctors = '⚠ ยังขาดคนคุม'
    elif any((e.get('need') or 0) > 0 for e in entries) and not staff:
        proctors += '  ⚠ ยังขาดคนคุม'
    dist, opener = names(dist_for(date, time, room)), short(openers.get(date, ''))
    dist_open = f'{dist} - {opener}' if dist and opener else (dist or opener)
    for e in entries:
        by_day.setdefault(date, []).append((time, e['code'], e['sec'], room, e['n'], proctors, dist_open))

wb = Workbook()
ws = wb.active
ws.title = 'ร่างตารางคุมสอบ'
thin = Side(style='thin', color='9AA7B0')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
HEAD = ['เวลา', 'วิชา', 'ตอน', 'ห้อง', 'นศ.', 'ผู้คุมสอบ\n(อาจารย์ - พนักงาน)', 'ผู้กระจายข้อสอบ - ผู้เปิดห้อง']
for col, w in zip('ABCDEFG', [14, 10, 8, 10, 6, 34, 30]):
    ws.column_dimensions[col].width = w

sem, year = '1', '2569'
m = re.match(r'(\d)/(\d{4})-(\w+)', data.get('round', ''))
if m:
    sem, year = m.group(1), m.group(2)
kind = 'ปลายภาค' if 'FINAL' in data.get('round', '') else 'กลางภาค'
ws['A1'] = f'(ร่าง) ตารางคุมสอบ{kind} ประจำภาคการศึกษาที่ {sem} ปีการศึกษา {year}'
ws['A1'].font = Font(bold=True, size=14)
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('A1:G1')
ws.row_dimensions[1].height = 24

r = 3
for date in sorted(by_day):
    ws.cell(r, 1, thai_day(date))
    ws.cell(r, 1).font = Font(bold=True, size=12, color='FFFFFF')
    ws.cell(r, 1).fill = PatternFill('solid', fgColor='1F3864')
    ws.cell(r, 1).alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 20
    r += 1
    for c, h in enumerate(HEAD, 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True, color='1F3864')
        cell.fill = PatternFill('solid', fgColor='DCE6F1')
        cell.alignment, cell.border = center, box
    ws.row_dimensions[r].height = 30
    r += 1
    for row in sorted(by_day[date], key=lambda x: (x[0], x[1], x[2], x[3])):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = left if c >= 6 else center
            cell.border = box
            if isinstance(v, str) and '⚠' in v:
                cell.font = Font(bold=True, color='9C0006')
                cell.fill = PatternFill('solid', fgColor='FFC7CE')
        r += 1

if data.get('problems'):
    r += 1
    ws.cell(r, 1, 'หมายเหตุ (ยังต้องตัดสินใจ)').font = Font(bold=True)
    for p in data['problems']:
        r += 1
        p = re.sub(r'\(\d/\d{4}-\w+\|(\d+)\|(\d+)\|\d+\)', r'วิชา \1 ตอน \2', p)
        ws.cell(r, 1, '• ' + p).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 32

ws.freeze_panes = 'A3'
ws.page_setup.orientation = 'portrait'
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
wb.save(out)
print('wrote', out, 'rows', sum(len(v) for v in by_day.values()))
