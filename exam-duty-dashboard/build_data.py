"""Extract exam-duty data from the optimizer workbook into data.json for the dashboard.

Usage: python3 build_data.py <optimize_FINAL.xlsx> [out.json]
"""
import json, re, sys, datetime as dt
import openpyxl

src = sys.argv[1]
out = sys.argv[2] if len(sys.argv) > 2 else 'data.json'
wb = openpyxl.load_workbook(src, data_only=True)

def rows(sheet, start=2):
    return [r for r in wb[sheet].iter_rows(min_row=start, values_only=True) if any(v is not None for v in r)]

summary = {r[0]: r[1] for r in rows('สรุป') if r[0]}
round_name = summary.get('รอบ', '')

# ---- rooms (one row per section x room) ----
rooms = {}
for r in rows('ห้องสอบ'):
    date, time, code, sec, lab, name, room, n, full, has_t, teacher, need, dist_help, rng, key = r[:15]
    rooms[f'{date}|{time}|{room}'] = rooms.get(f'{date}|{time}|{room}', []) + [{
        'code': code, 'sec': sec, 'name': name, 'room': room, 'n': n, 'full': full,
        'teacher': teacher, 'need': need, 'distHelp': dist_help, 'range': rng, 'key': key,
    }]

# ---- duties ----
ROLE = {'คุมสอบ': 'proctor', 'จ่ายข้อสอบ': 'dist', 'เปิดห้อง': 'opener', 'อาจารย์คุมวิชาตนเอง': 'own'}
duty_rows = rows('เวรคุมสอบ')
people = {}
duties = []
for date, dow, time, p, role, name, email, key, room in duty_rows:
    pid = email
    kind = 'teacher' if role == 'อาจารย์คุมวิชาตนเอง' else 'staff'
    if pid not in people:
        people[pid] = {'id': pid, 'name': name, 'kind': kind}
    duties.append({'pid': pid, 'date': date, 'time': time, 'p': p, 'role': ROLE[role], 'room': room})

# ---- staff load sheet (adds people with zero duties + notes) ----
for r in rows('ภาระต่อคน'):
    email, name = r[0], r[1]
    note = r[10] if isinstance(r[10], str) and r[10] != 'ใช่' else None
    if email == 'ketdao.l@cmu.ac.th':  # resigned, not in this round
        continue
    p = people.setdefault(email, {'id': email, 'name': name, 'kind': 'staff'})
    p['kind'] = 'staff'
    p['canProctor'] = r[11] == 'ใช่'
    p['inDistPool'] = r[10] == 'ใช่'
    if note:
        p['note'] = note

# ---- cumulative stacks (weekday x period) before / after this round ----
ws = list(wb['สแตครายวันต่อคน'].iter_rows(values_only=True))
cols = None
section = None
for r in ws:
    if r[0] and r[1] is None:
        section = 'before' if r[0].startswith('ก่อน') else 'after'
        continue
    if r[0] == 'อีเมล':
        cols = list(r[2:23])
        continue
    if r[0] and section and r[0] in people:
        people[r[0]][section] = [int(v or 0) for v in r[2:23]]
stack_cols = cols

# ---- slots & days ----
slots = sorted({(d['date'], d['time'], d['p']) for d in duties})
dates = sorted({s[0] for s in slots})
d0 = dt.date.fromisoformat(dates[0]); d1 = dt.date.fromisoformat(dates[-1])
all_days = [(d0 + dt.timedelta(i)).isoformat() for i in range((d1 - d0).days + 1)]

problems = [r[1] for r in rows('ปัญหา') if r[1]]
unavail = [{'pid': r[0], 'date': r[2], 'reason': r[4]} for r in rows('วันไม่ว่าง-เวรระบบอื่น') if r[0]]

# emails are ids only; the page never renders them. Replace with opaque ids.
id_map = {pid: f'p{i}' for i, pid in enumerate(sorted(people))}
for p in people.values():
    p['id'] = id_map[p['id']]
for d in duties:
    d['pid'] = id_map[d['pid']]
for u in unavail:
    u['pid'] = id_map.get(u['pid'])

data = {
    'round': round_name,
    'generated': dt.date.today().isoformat(),
    'summary': {
        'rooms': summary.get('รายการห้องที่ใช้ (1 ตอนอาจใช้หลายห้อง)'),
        'slots': summary.get('ช่วงสอบทั้งหมด (วัน x เวลา)'),
        'duties': summary.get('เวรทั้งหมด'),
        'cap': summary.get('เพดานเวรต่อคนในรอบนี้ (base cap)'),
    },
    'days': all_days,
    'slots': [{'date': a, 'time': b, 'p': c} for a, b, c in slots],
    'stackCols': stack_cols,
    'people': sorted(people.values(), key=lambda p: (p['kind'] != 'staff', p['name'])),
    'duties': duties,
    'rooms': rooms,
    'problems': problems,
    'unavail': [u for u in unavail if u['pid']],
}
with open(out, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
print(f'wrote {out}: {len(people)} people, {len(duties)} duties, {len(slots)} slots, {len(rooms)} room-slots')
